import os
import uuid
import PyPDF2
import pdfplumber
import docx
import docx2txt
import pandas as pd
import zipfile
import rarfile
import chardet
import mimetypes
from typing import Dict, Any, Optional, Tuple, List
import logging
import re
import io
from pathlib import Path
from werkzeug.utils import secure_filename
from werkzeug.datastructures import FileStorage

# Try to import additional libraries for better text extraction
try:
    import fitz  # PyMuPDF for better PDF handling
    HAS_PYMUPDF = True
except ImportError:
    HAS_PYMUPDF = False

try:
    import textract  # For various file formats
    HAS_TEXTRACT = True
except ImportError:
    HAS_TEXTRACT = False

try:
    from pdfminer.high_level import extract_text as pdfminer_extract
    from pdfminer.layout import LAParams
    HAS_PDFMINER = True
except ImportError:
    HAS_PDFMINER = False

try:
    import easyocr
    HAS_OCR = True
except ImportError:
    HAS_OCR = False

try:
    from PIL import Image
    import pytesseract
    HAS_TESSERACT = True
except ImportError:
    HAS_TESSERACT = False

logger = logging.getLogger(__name__)

class ResumeParsingService:
    """Enhanced service for parsing and extracting text from resume files with robust handling."""
    
    def __init__(self, upload_folder: str):
        self.upload_folder = upload_folder
        # Expanded list of supported file extensions
        self.allowed_extensions = {
            'pdf', 'doc', 'docx', 'txt', 'rtf', 'odt', 'pages',
            'xls', 'xlsx', 'csv', 'html', 'htm', 'xml',
            'jpg', 'jpeg', 'png', 'tiff', 'bmp', 'gif',  # Image formats for OCR
            'zip', 'rar'  # Archive formats
        }
        
        # Text extraction preferences (ordered by reliability)
        self.pdf_extractors = ['pdfplumber', 'pymupdf', 'pdfminer', 'pypdf2', 'textract']
        self.ocr_enabled = HAS_OCR or HAS_TESSERACT
        
        # Initialize OCR if available
        self.ocr_reader = None
        if HAS_OCR:
            try:
                self.ocr_reader = easyocr.Reader(['en'])
            except Exception as e:
                logger.warning(f"Could not initialize EasyOCR: {e}")
        
        # Ensure upload directory exists
        os.makedirs(upload_folder, exist_ok=True)
    
    def is_allowed_file(self, filename: str) -> bool:
        """Check if file extension is allowed."""
        if not filename or '.' not in filename:
            return False
        
        extension = filename.rsplit('.', 1)[1].lower()
        return extension in self.allowed_extensions
    
    def detect_file_type(self, file_path: str) -> str:
        """Detect actual file type using MIME type detection."""
        try:
            mime_type, _ = mimetypes.guess_type(file_path)
            
            # Read file signature for more accurate detection
            with open(file_path, 'rb') as f:
                signature = f.read(16)
            
            # PDF signatures
            if signature.startswith(b'%PDF'):
                return 'pdf'
            
            # Microsoft Office signatures
            elif signature.startswith(b'PK\x03\x04') or signature.startswith(b'PK\x05\x06'):
                # Could be docx, xlsx, or zip
                if mime_type and 'word' in mime_type:
                    return 'docx'
                elif mime_type and 'excel' in mime_type:
                    return 'xlsx'
                elif file_path.lower().endswith(('.docx', '.doc')):
                    return 'docx'
                elif file_path.lower().endswith(('.xlsx', '.xls')):
                    return 'xlsx'
                else:
                    return 'zip'
            
            # RTF signature
            elif signature.startswith(b'{\\rtf'):
                return 'rtf'
            
            # Image signatures
            elif signature.startswith(b'\xff\xd8\xff'):
                return 'jpeg'
            elif signature.startswith(b'\x89PNG'):
                return 'png'
            elif signature.startswith(b'GIF'):
                return 'gif'
            elif signature.startswith(b'BM'):
                return 'bmp'
            
            # Fallback to extension
            extension = file_path.rsplit('.', 1)[1].lower() if '.' in file_path else 'unknown'
            return extension
            
        except Exception as e:
            logger.warning(f"Could not detect file type for {file_path}: {e}")
            return file_path.rsplit('.', 1)[1].lower() if '.' in file_path else 'unknown'
    
    def save_file(self, file: FileStorage) -> Tuple[str, str, Dict[str, Any]]:
        """
        Save uploaded file and return file info with enhanced validation.
        
        Args:
            file: Uploaded file object
            
        Returns:
            Tuple of (filename, file_path, metadata)
        """
        if not file or not file.filename:
            raise ValueError("No file provided")
        
        if not self.is_allowed_file(file.filename):
            raise ValueError(f"File type not allowed. Allowed types: {self.allowed_extensions}")
        
        # Generate unique filename
        original_filename = secure_filename(file.filename)
        file_extension = original_filename.rsplit('.', 1)[1].lower() if '.' in original_filename else 'txt'
        unique_filename = f"{uuid.uuid4()}.{file_extension}"
        file_path = os.path.join(self.upload_folder, unique_filename)
        
        # Save file
        file.save(file_path)
        
        # Validate file was saved correctly
        if not os.path.exists(file_path):
            raise Exception("File save failed - file not found after save")
        
        # Get file metadata
        file_size = os.path.getsize(file_path)
        
        # Detect actual file type
        detected_type = self.detect_file_type(file_path)
        
        metadata = {
            'original_filename': original_filename,
            'filename': unique_filename,
            'file_path': file_path,
            'file_size': file_size,
            'file_type': file_extension,
            'detected_type': detected_type,
            'mime_type': mimetypes.guess_type(file_path)[0]
        }
        
        logger.info(f"File saved: {unique_filename} ({file_size} bytes, detected: {detected_type})")
        return unique_filename, file_path, metadata
    
    def extract_text(self, file_path: str) -> str:
        """
        Enhanced text extraction with multiple fallback methods.
        
        Args:
            file_path: Path to the resume file
            
        Returns:
            Extracted text content
        """
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File not found: {file_path}")
        
        # Detect actual file type
        detected_type = self.detect_file_type(file_path)
        file_extension = file_path.rsplit('.', 1)[1].lower() if '.' in file_path else detected_type
        
        logger.info(f"Extracting text from {file_path} (detected: {detected_type}, extension: {file_extension})")
        
        try:
            # Handle different file types with multiple extraction methods
            if detected_type == 'pdf' or file_extension == 'pdf':
                return self._extract_from_pdf_robust(file_path)
            
            elif detected_type in ['docx', 'doc'] or file_extension in ['docx', 'doc']:
                return self._extract_from_word_robust(file_path)
            
            elif detected_type == 'txt' or file_extension == 'txt':
                return self._extract_from_text_robust(file_path)
            
            elif detected_type == 'rtf' or file_extension == 'rtf':
                return self._extract_from_rtf(file_path)
            
            elif detected_type in ['xlsx', 'xls', 'csv'] or file_extension in ['xlsx', 'xls', 'csv']:
                return self._extract_from_spreadsheet(file_path)
            
            elif detected_type in ['html', 'htm', 'xml'] or file_extension in ['html', 'htm', 'xml']:
                return self._extract_from_html(file_path)
            
            elif detected_type in ['jpg', 'jpeg', 'png', 'tiff', 'bmp', 'gif'] or file_extension in ['jpg', 'jpeg', 'png', 'tiff', 'bmp', 'gif']:
                return self._extract_from_image(file_path)
            
            elif detected_type in ['zip', 'rar'] or file_extension in ['zip', 'rar']:
                return self._extract_from_archive(file_path)
            
            elif detected_type == 'odt' or file_extension == 'odt':
                return self._extract_from_odt(file_path)
            
            else:
                # Try universal text extraction methods
                return self._extract_with_fallback_methods(file_path)
                
        except Exception as e:
            logger.error(f"Primary text extraction failed for {file_path}: {str(e)}")
            # Final fallback - try all available methods
            return self._extract_with_all_methods(file_path)
    
    def _extract_from_pdf_robust(self, file_path: str) -> str:
        """Enhanced PDF text extraction with multiple methods."""
        texts = []
        methods_tried = []
        
        # Method 1: pdfplumber (best for tables and structured content)
        try:
            import pdfplumber
            with pdfplumber.open(file_path) as pdf:
                text = ""
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n"
                if text.strip():
                    texts.append(('pdfplumber', text.strip()))
                    methods_tried.append('pdfplumber')
        except Exception as e:
            logger.debug(f"pdfplumber failed: {e}")
        
        # Method 2: PyMuPDF (excellent for complex PDFs)
        if HAS_PYMUPDF:
            try:
                doc = fitz.open(file_path)
                text = ""
                for page in doc:
                    text += page.get_text() + "\n"
                doc.close()
                if text.strip():
                    texts.append(('pymupdf', text.strip()))
                    methods_tried.append('pymupdf')
            except Exception as e:
                logger.debug(f"PyMuPDF failed: {e}")
        
        # Method 3: pdfminer (good for text-heavy PDFs)
        if HAS_PDFMINER:
            try:
                laparams = LAParams(
                    line_margin=0.5,
                    word_margin=0.1,
                    char_margin=2.0,
                    boxes_flow=0.5,
                    all_texts=False
                )
                text = pdfminer_extract(file_path, laparams=laparams)
                if text.strip():
                    texts.append(('pdfminer', text.strip()))
                    methods_tried.append('pdfminer')
            except Exception as e:
                logger.debug(f"pdfminer failed: {e}")
        
        # Method 4: PyPDF2 (basic but reliable)
        try:
            with open(file_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                text = ""
                for page in pdf_reader.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n"
                if text.strip():
                    texts.append(('pypdf2', text.strip()))
                    methods_tried.append('pypdf2')
        except Exception as e:
            logger.debug(f"PyPDF2 failed: {e}")
        
        # Method 5: textract (universal extractor)
        if HAS_TEXTRACT:
            try:
                text = textract.process(file_path).decode('utf-8')
                if text.strip():
                    texts.append(('textract', text.strip()))
                    methods_tried.append('textract')
            except Exception as e:
                logger.debug(f"textract failed: {e}")
        
        # Method 6: OCR for scanned PDFs
        if not texts and self.ocr_enabled:
            try:
                text = self._extract_pdf_with_ocr(file_path)
                if text.strip():
                    texts.append(('ocr', text.strip()))
                    methods_tried.append('ocr')
            except Exception as e:
                logger.debug(f"OCR failed: {e}")
        
        # Choose the best result (longest text that makes sense)
        if texts:
            best_text = self._choose_best_extraction(texts)
            logger.info(f"PDF extraction successful using methods: {methods_tried}")
            return best_text
        
        raise Exception("All PDF extraction methods failed")
    
    def _extract_from_word_robust(self, file_path: str) -> str:
        """Enhanced Word document extraction."""
        texts = []
        methods_tried = []
        
        # Method 1: python-docx (best for .docx)
        try:
            doc = docx.Document(file_path)
            text = []
            
            # Extract paragraphs
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text.append(paragraph.text)
            
            # Extract tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = []
                    for cell in row.cells:
                        if cell.text.strip():
                            row_text.append(cell.text.strip())
                    if row_text:
                        text.append(" | ".join(row_text))
            
            full_text = "\n".join(text)
            if full_text.strip():
                texts.append(('python-docx', full_text.strip()))
                methods_tried.append('python-docx')
        except Exception as e:
            logger.debug(f"python-docx failed: {e}")
        
        # Method 2: docx2txt (fallback for .docx)
        try:
            import docx2txt
            text = docx2txt.process(file_path)
            if text.strip():
                texts.append(('docx2txt', text.strip()))
                methods_tried.append('docx2txt')
        except Exception as e:
            logger.debug(f"docx2txt failed: {e}")
        
        # Method 3: textract (universal)
        if HAS_TEXTRACT:
            try:
                text = textract.process(file_path).decode('utf-8')
                if text.strip():
                    texts.append(('textract', text.strip()))
                    methods_tried.append('textract')
            except Exception as e:
                logger.debug(f"textract failed: {e}")
        
        if texts:
            best_text = self._choose_best_extraction(texts)
            logger.info(f"Word extraction successful using methods: {methods_tried}")
            return best_text
        
        raise Exception("All Word extraction methods failed")
    
    def _extract_from_text_robust(self, file_path: str) -> str:
        """Enhanced text file extraction with encoding detection."""
        # Try to detect encoding
        try:
            with open(file_path, 'rb') as file:
                raw_data = file.read()
                encoding_result = chardet.detect(raw_data)
                encoding = encoding_result['encoding'] if encoding_result['confidence'] > 0.7 else 'utf-8'
        except:
            encoding = 'utf-8'
        
        # List of encodings to try
        encodings_to_try = [encoding, 'utf-8', 'latin-1', 'cp1252', 'iso-8859-1', 'ascii']
        
        for enc in encodings_to_try:
            try:
                with open(file_path, 'r', encoding=enc) as file:
                    text = file.read().strip()
                    if text:
                        logger.info(f"Text extraction successful using encoding: {enc}")
                        return text
            except Exception as e:
                logger.debug(f"Encoding {enc} failed: {e}")
                continue
        
        raise Exception("All text encoding methods failed")
    
    def _extract_from_rtf(self, file_path: str) -> str:
        """Extract text from RTF files."""
        try:
            # Try textract first
            if HAS_TEXTRACT:
                text = textract.process(file_path).decode('utf-8')
                return text.strip()
        except Exception as e:
            logger.debug(f"textract RTF failed: {e}")
        
        # Basic RTF parsing
        try:
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as file:
                content = file.read()
                # Remove RTF control words (basic parsing)
                text = re.sub(r'\\[a-z]+\d*\s?', '', content)
                text = re.sub(r'[{}]', '', text)
                return text.strip()
        except Exception as e:
            raise Exception(f"RTF extraction failed: {str(e)}")
    
    def _extract_from_spreadsheet(self, file_path: str) -> str:
        """Extract text from Excel/CSV files."""
        try:
            # Determine file type
            file_extension = file_path.rsplit('.', 1)[1].lower()
            
            if file_extension == 'csv':
                df = pd.read_csv(file_path, encoding='utf-8')
            else:
                df = pd.read_excel(file_path, sheet_name=None)  # Read all sheets
                
                # Combine all sheets
                if isinstance(df, dict):
                    combined_df = pd.DataFrame()
                    for sheet_name, sheet_df in df.items():
                        combined_df = pd.concat([combined_df, sheet_df], ignore_index=True)
                    df = combined_df
            
            # Convert to text
            text_lines = []
            for column in df.columns:
                text_lines.append(f"{column}: {', '.join(df[column].astype(str).dropna().unique())}")
            
            return '\n'.join(text_lines)
            
        except Exception as e:
            raise Exception(f"Spreadsheet extraction failed: {str(e)}")
    
    def _extract_from_html(self, file_path: str) -> str:
        """Extract text from HTML/XML files."""
        try:
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as file:
                content = file.read()
            
            # Remove HTML tags
            text = re.sub(r'<[^>]+>', '', content)
            # Clean up whitespace
            text = re.sub(r'\s+', ' ', text)
            return text.strip()
            
        except Exception as e:
            raise Exception(f"HTML extraction failed: {str(e)}")
    
    def _extract_from_image(self, file_path: str) -> str:
        """Extract text from images using OCR."""
        if not self.ocr_enabled:
            raise Exception("OCR not available for image text extraction")
        
        try:
            # Try EasyOCR first
            if self.ocr_reader:
                result = self.ocr_reader.readtext(file_path)
                text = ' '.join([item[1] for item in result])
                if text.strip():
                    return text.strip()
            
            # Try Tesseract as fallback
            if HAS_TESSERACT:
                image = Image.open(file_path)
                text = pytesseract.image_to_string(image)
                return text.strip()
            
            raise Exception("No OCR method available")
            
        except Exception as e:
            raise Exception(f"Image OCR extraction failed: {str(e)}")
    
    def _extract_from_archive(self, file_path: str) -> str:
        """Extract text from archive files (ZIP, RAR)."""
        import tempfile
        import shutil
        
        with tempfile.TemporaryDirectory() as temp_dir:
            try:
                # Extract archive
                if file_path.lower().endswith('.zip'):
                    with zipfile.ZipFile(file_path, 'r') as zip_ref:
                        zip_ref.extractall(temp_dir)
                elif file_path.lower().endswith('.rar'):
                    try:
                        import rarfile
                        with rarfile.RarFile(file_path, 'r') as rar_ref:
                            rar_ref.extractall(temp_dir)
                    except ImportError:
                        raise Exception("rarfile library not available")
                
                # Extract text from all files in archive
                all_text = []
                for root, dirs, files in os.walk(temp_dir):
                    for file in files:
                        file_path_inner = os.path.join(root, file)
                        if self.is_allowed_file(file):
                            try:
                                text = self.extract_text(file_path_inner)
                                if text.strip():
                                    all_text.append(f"--- File: {file} ---\n{text}")
                            except Exception as e:
                                logger.debug(f"Failed to extract from {file}: {e}")
                
                return '\n\n'.join(all_text)
                
            except Exception as e:
                raise Exception(f"Archive extraction failed: {str(e)}")
    
    def _extract_from_odt(self, file_path: str) -> str:
        """Extract text from OpenDocument Text files."""
        try:
            if HAS_TEXTRACT:
                text = textract.process(file_path).decode('utf-8')
                return text.strip()
        except Exception as e:
            logger.debug(f"textract ODT failed: {e}")
        
        # Try as ZIP (ODT is actually a ZIP file)
        try:
            import xml.etree.ElementTree as ET
            with zipfile.ZipFile(file_path, 'r') as odt_zip:
                content_xml = odt_zip.read('content.xml')
                root = ET.fromstring(content_xml)
                
                # Extract text from XML
                text_parts = []
                for elem in root.iter():
                    if elem.text:
                        text_parts.append(elem.text)
                
                return ' '.join(text_parts).strip()
                
        except Exception as e:
            raise Exception(f"ODT extraction failed: {str(e)}")
    
    def _extract_with_fallback_methods(self, file_path: str) -> str:
        """Try universal extraction methods."""
        methods = []
        
        # Try textract if available
        if HAS_TEXTRACT:
            try:
                text = textract.process(file_path).decode('utf-8')
                if text.strip():
                    return text.strip()
            except Exception as e:
                logger.debug(f"textract fallback failed: {e}")
        
        # Try reading as text with various encodings
        try:
            return self._extract_from_text_robust(file_path)
        except Exception as e:
            logger.debug(f"text fallback failed: {e}")
        
        raise Exception("All fallback extraction methods failed")
    
    def _extract_with_all_methods(self, file_path: str) -> str:
        """Final fallback - try every available method."""
        logger.warning(f"Attempting final fallback extraction for {file_path}")
        
        # Try treating as different file types
        methods_to_try = [
            ('pdf', self._extract_from_pdf_robust),
            ('word', self._extract_from_word_robust),
            ('text', self._extract_from_text_robust),
            ('fallback', self._extract_with_fallback_methods)
        ]
        
        for method_name, method_func in methods_to_try:
            try:
                text = method_func(file_path)
                if text.strip():
                    logger.info(f"Final fallback successful using {method_name} method")
                    return text.strip()
            except Exception as e:
                logger.debug(f"Final fallback {method_name} failed: {e}")
        
        # Last resort - return file path as text
        return f"Could not extract text from file: {os.path.basename(file_path)}"
    
    def _extract_pdf_with_ocr(self, file_path: str) -> str:
        """Extract text from PDF using OCR for scanned documents."""
        if not HAS_PYMUPDF or not self.ocr_enabled:
            raise Exception("OCR requirements not met")
        
        try:
            doc = fitz.open(file_path)
            full_text = []
            
            for page_num in range(len(doc)):
                page = doc.load_page(page_num)
                
                # First try regular text extraction
                text = page.get_text()
                if text.strip():
                    full_text.append(text)
                else:
                    # Convert page to image and OCR
                    pix = page.get_pixmap()
                    img_data = pix.tobytes("png")
                    
                    if self.ocr_reader:
                        # Use EasyOCR
                        result = self.ocr_reader.readtext(img_data)
                        ocr_text = ' '.join([item[1] for item in result])
                        full_text.append(ocr_text)
                    elif HAS_TESSERACT:
                        # Use Tesseract
                        image = Image.open(io.BytesIO(img_data))
                        ocr_text = pytesseract.image_to_string(image)
                        full_text.append(ocr_text)
            
            doc.close()
            return '\n'.join(full_text).strip()
            
        except Exception as e:
            raise Exception(f"PDF OCR extraction failed: {str(e)}")
    
    def _choose_best_extraction(self, texts: List[Tuple[str, str]]) -> str:
        """Choose the best text extraction result."""
        if not texts:
            raise Exception("No extraction results to choose from")
        
        if len(texts) == 1:
            return texts[0][1]
        
        # Score each result based on length and content quality
        scored_results = []
        for method, text in texts:
            score = 0
            
            # Length score (longer is generally better)
            score += len(text) * 0.1
            
            # Content quality indicators
            if '@' in text:  # Contains email
                score += 100
            if re.search(r'\d{3}[-.\s]?\d{3}[-.\s]?\d{4}', text):  # Contains phone
                score += 100
            if re.search(r'\b(experience|education|skills|projects)\b', text, re.IGNORECASE):
                score += 200
            
            # Penalize results with too many special characters (usually extraction errors)
            special_char_ratio = len(re.findall(r'[^\w\s]', text)) / len(text) if text else 1
            if special_char_ratio > 0.3:
                score -= 500
            
            scored_results.append((score, method, text))
        
        # Return the highest scoring result
        best_result = max(scored_results, key=lambda x: x[0])
        logger.info(f"Chose {best_result[1]} method (score: {best_result[0]:.1f})")
        return best_result[2]
    
    def parse_resume_structure(self, text: str) -> Dict[str, Any]:
        """
        Enhanced resume structure parsing with better section detection.
        
        Args:
            text: Raw resume text
            
        Returns:
            Dictionary with structured resume data
        """
        # Clean and normalize text
        cleaned_text = self._clean_text(text)
        
        structure = {
            'contact_info': self._extract_contact_info(cleaned_text),
            'summary': self._extract_summary(cleaned_text),
            'experience': self._extract_experience_enhanced(cleaned_text),
            'education': self._extract_education_enhanced(cleaned_text),
            'skills': self._extract_skills_enhanced(cleaned_text),
            'certifications': self._extract_certifications_enhanced(cleaned_text),
            'projects': self._extract_projects(cleaned_text),
            'languages': self._extract_languages(cleaned_text),
            'achievements': self._extract_achievements(cleaned_text),
            'additional_sections': self._extract_additional_sections(cleaned_text),
            'metadata': {
                'total_lines': len(cleaned_text.split('\n')),
                'total_words': len(cleaned_text.split()),
                'total_characters': len(cleaned_text),
                'sections_found': []
            }
        }
        
        # Track which sections were successfully extracted
        for key, value in structure.items():
            if key != 'metadata' and value:
                if isinstance(value, (list, dict)) and value:
                    structure['metadata']['sections_found'].append(key)
                elif isinstance(value, str) and value.strip():
                    structure['metadata']['sections_found'].append(key)
        
        return structure
    
    def _clean_text(self, text: str) -> str:
        """Clean and normalize extracted text."""
        if not text:
            return ""
        
        # Remove excessive whitespace
        text = re.sub(r'\s+', ' ', text)
        
        # Remove non-printable characters except newlines and tabs
        text = re.sub(r'[^\x20-\x7E\n\t]', '', text)
        
        # Normalize line breaks
        text = re.sub(r'\r\n', '\n', text)
        text = re.sub(r'\r', '\n', text)
        
        # Remove excessive blank lines
        text = re.sub(r'\n\s*\n\s*\n', '\n\n', text)
        
        return text.strip()
    
    def _extract_contact_info(self, text: str) -> Dict[str, Any]:
        """Enhanced contact information extraction."""
        contact_info = {
            'emails': [],
            'phones': [],
            'linkedin': None,
            'github': None,
            'website': None,
            'address': None,
            'name': None
        }
        
        # Email extraction
        email_pattern = r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b'
        contact_info['emails'] = list(set(re.findall(email_pattern, text)))
        
        # Phone extraction (multiple formats)
        phone_patterns = [
            r'(\+?\d{1,3}[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}',
            r'\+?\d{1,3}[-.\s]?\d{3,4}[-.\s]?\d{3,4}[-.\s]?\d{3,4}',
            r'\b\d{10}\b'
        ]
        
        phones = []
        for pattern in phone_patterns:
            phones.extend(re.findall(pattern, text))
        contact_info['phones'] = list(set([phone for phone in phones if phone]))
        
        # LinkedIn profile
        linkedin_match = re.search(r'linkedin\.com/in/[\w-]+', text, re.IGNORECASE)
        if linkedin_match:
            contact_info['linkedin'] = linkedin_match.group()
        
        # GitHub profile
        github_match = re.search(r'github\.com/[\w-]+', text, re.IGNORECASE)
        if github_match:
            contact_info['github'] = github_match.group()
        
        # Website/Portfolio
        website_match = re.search(r'https?://[\w\.-]+\.[a-zA-Z]{2,}', text)
        if website_match:
            contact_info['website'] = website_match.group()
        
        # Name extraction (first few lines, common patterns)
        lines = text.split('\n')[:5]
        name_patterns = [
            r'^([A-Z][a-z]+ [A-Z][a-z]+(?:\s[A-Z][a-z]+)?)',
            r'Name:?\s*([A-Z][a-z]+ [A-Z][a-z]+(?:\s[A-Z][a-z]+)?)'
        ]
        
        for line in lines:
            for pattern in name_patterns:
                match = re.search(pattern, line)
                if match:
                    contact_info['name'] = match.group(1)
                    break
            if contact_info['name']:
                break
        
        return contact_info
    
    def _extract_experience_enhanced(self, text: str) -> List[Dict[str, Any]]:
        """Enhanced work experience extraction."""
        experience_sections = []
        
        # Look for experience section
        exp_keywords = ['experience', 'employment', 'work history', 'professional experience', 'career']
        sections = self._split_into_sections(text)
        
        for section_title, section_content in sections:
            if any(keyword in section_title.lower() for keyword in exp_keywords):
                # Parse individual experiences
                experiences = self._parse_experience_entries(section_content)
                experience_sections.extend(experiences)
        
        return experience_sections
    
    def _extract_education_enhanced(self, text: str) -> List[Dict[str, Any]]:
        """Enhanced education extraction."""
        education_entries = []
        
        # Look for education section
        edu_keywords = ['education', 'academic', 'qualification', 'degree']
        sections = self._split_into_sections(text)
        
        for section_title, section_content in sections:
            if any(keyword in section_title.lower() for keyword in edu_keywords):
                # Parse individual education entries
                entries = self._parse_education_entries(section_content)
                education_entries.extend(entries)
        
        return education_entries
    
    def _extract_skills_enhanced(self, text: str) -> Dict[str, List[str]]:
        """Enhanced skills extraction with categorization."""
        skills = {
            'technical': [],
            'programming': [],
            'tools': [],
            'soft_skills': [],
            'languages': [],
            'other': []
        }
        
        # Look for skills section
        skill_keywords = ['skills', 'competencies', 'expertise', 'technologies', 'technical skills']
        sections = self._split_into_sections(text)
        
        for section_title, section_content in sections:
            if any(keyword in section_title.lower() for keyword in skill_keywords):
                parsed_skills = self._categorize_skills(section_content)
                for category, skill_list in parsed_skills.items():
                    skills[category].extend(skill_list)
        
        # Remove duplicates
        for category in skills:
            skills[category] = list(set(skills[category]))
        
        return skills
    
    def _extract_projects(self, text: str) -> List[Dict[str, Any]]:
        """Extract project information."""
        projects = []
        
        project_keywords = ['projects', 'portfolio', 'work samples', 'achievements']
        sections = self._split_into_sections(text)
        
        for section_title, section_content in sections:
            if any(keyword in section_title.lower() for keyword in project_keywords):
                project_entries = self._parse_project_entries(section_content)
                projects.extend(project_entries)
        
        return projects
    
    def _extract_languages(self, text: str) -> List[Dict[str, str]]:
        """Extract language proficiency information."""
        languages = []
        
        # Common language indicators
        language_pattern = r'(Hindi|English|Tamil|Telugu|Marathi|Bengali|Gujarati|Kannada|Malayalam|Punjabi|Urdu|Sanskrit|French|German|Spanish|Chinese|Japanese|Korean|Arabic)\s*[-:,]?\s*(Native|Fluent|Advanced|Intermediate|Basic|Beginner|Professional)?'
        
        matches = re.findall(language_pattern, text, re.IGNORECASE)
        for language, proficiency in matches:
            languages.append({
                'language': language.title(),
                'proficiency': proficiency.title() if proficiency else 'Unknown'
            })
        
        return languages
    
    def _extract_achievements(self, text: str) -> List[str]:
        """Extract achievements and awards."""
        achievements = []
        
        achievement_keywords = ['achievement', 'award', 'recognition', 'honor', 'certification', 'accomplishment']
        sections = self._split_into_sections(text)
        
        for section_title, section_content in sections:
            if any(keyword in section_title.lower() for keyword in achievement_keywords):
                # Split into individual achievements
                lines = section_content.split('\n')
                for line in lines:
                    line = line.strip()
                    if line and len(line) > 10:  # Filter out very short lines
                        achievements.append(line)
        
        return achievements
    
    def _split_into_sections(self, text: str) -> List[Tuple[str, str]]:
        """Split resume text into sections based on headers."""
        sections = []
        lines = text.split('\n')
        
        current_section = ""
        current_content = []
        
        for line in lines:
            line = line.strip()
            if not line:
                continue
            
            # Check if line is a section header (all caps, or ends with colon, etc.)
            if (line.isupper() or 
                line.endswith(':') or 
                re.match(r'^[A-Z][A-Z\s&/]+$', line) or
                any(keyword in line.lower() for keyword in ['experience', 'education', 'skills', 'projects', 'certifications'])):
                
                # Save previous section
                if current_section and current_content:
                    sections.append((current_section, '\n'.join(current_content)))
                
                # Start new section
                current_section = line.replace(':', '').strip()
                current_content = []
            else:
                current_content.append(line)
        
        # Add last section
        if current_section and current_content:
            sections.append((current_section, '\n'.join(current_content)))
        
        return sections
    
    def _parse_experience_entries(self, content: str) -> List[Dict[str, Any]]:
        """Parse individual work experience entries."""
        experiences = []
        
        # Split by potential job entries (look for company names, job titles)
        job_pattern = r'([A-Z][a-zA-Z\s&,.-]+(?:Company|Corp|Inc|Ltd|LLC|Technologies|Systems|Solutions|Services|Group|Pvt|Private|Limited))'
        
        entries = re.split(job_pattern, content)
        
        for i in range(1, len(entries), 2):
            if i + 1 < len(entries):
                company = entries[i].strip()
                details = entries[i + 1].strip()
                
                experience = {
                    'company': company,
                    'title': self._extract_job_title(details),
                    'duration': self._extract_duration(details),
                    'description': details,
                    'location': self._extract_location(details)
                }
                experiences.append(experience)
        
        return experiences
    
    def _parse_education_entries(self, content: str) -> List[Dict[str, Any]]:
        """Parse individual education entries."""
        education = []
        
        # Look for degree patterns
        degree_pattern = r'(Bachelor|Master|PhD|Doctorate|Diploma|Certificate|B\.?Tech|M\.?Tech|B\.?E|M\.?E|B\.?Sc|M\.?Sc|MBA|BBA|B\.?Com|M\.?Com)'
        
        lines = content.split('\n')
        for line in lines:
            if re.search(degree_pattern, line, re.IGNORECASE):
                education.append({
                    'degree': line.strip(),
                    'institution': self._extract_institution(line),
                    'year': self._extract_year(line),
                    'grade': self._extract_grade(line)
                })
        
        return education
    
    def _categorize_skills(self, content: str) -> Dict[str, List[str]]:
        """Categorize skills into different types."""
        categorized = {
            'technical': [],
            'programming': [],
            'tools': [],
            'soft_skills': [],
            'other': []
        }
        
        # Programming languages
        programming_langs = ['Python', 'Java', 'JavaScript', 'C++', 'C#', 'PHP', 'Ruby', 'Go', 'Swift', 'Kotlin']
        
        # Tools and technologies
        tools = ['Docker', 'Kubernetes', 'Git', 'Jenkins', 'AWS', 'Azure', 'Linux', 'Windows', 'MySQL', 'PostgreSQL']
        
        # Soft skills
        soft_skills = ['Leadership', 'Communication', 'Teamwork', 'Problem Solving', 'Time Management']
        
        # Extract skills from content
        skill_items = re.split(r'[,;\n\|•]', content)
        
        for item in skill_items:
            item = item.strip()
            if not item:
                continue
            
            if any(lang.lower() in item.lower() for lang in programming_langs):
                categorized['programming'].append(item)
            elif any(tool.lower() in item.lower() for tool in tools):
                categorized['tools'].append(item)
            elif any(skill.lower() in item.lower() for skill in soft_skills):
                categorized['soft_skills'].append(item)
            else:
                categorized['other'].append(item)
        
        return categorized
    
    def _parse_project_entries(self, content: str) -> List[Dict[str, Any]]:
        """Parse project entries."""
        projects = []
        
        # Split by project indicators
        project_lines = content.split('\n')
        current_project = None
        
        for line in project_lines:
            line = line.strip()
            if not line:
                continue
            
            # Check if it's a project title (usually starts with capital letter, contains project-like words)
            if (line[0].isupper() and 
                any(word in line.lower() for word in ['project', 'system', 'application', 'website', 'portal'])):
                
                if current_project:
                    projects.append(current_project)
                
                current_project = {
                    'name': line,
                    'description': '',
                    'technologies': []
                }
            elif current_project:
                current_project['description'] += line + ' '
        
        if current_project:
            projects.append(current_project)
        
        return projects
    
    def _extract_job_title(self, text: str) -> str:
        """Extract job title from text."""
        # Common job title patterns
        title_patterns = [
            r'(Senior|Junior|Lead|Principal|Chief|Assistant|Associate)?\s*(Developer|Engineer|Manager|Analyst|Consultant|Specialist|Director|Designer|Architect)',
            r'(Software|Web|Mobile|Full Stack|Frontend|Backend|Data|System|Network|Security)\s+(Developer|Engineer|Analyst)',
        ]
        
        for pattern in title_patterns:
            match = re.search(pattern, text, re.IGNORECASE)
            if match:
                return match.group().strip()
        
        return ""
    
    def _extract_duration(self, text: str) -> str:
        """Extract job duration from text."""
        duration_patterns = [
            r'(\d{4})\s*-\s*(\d{4}|Present|Current)',
            r'(\w+ \d{4})\s*-\s*(\w+ \d{4}|Present|Current)',
            r'(\d+)\s+(years?|months?)',
        ]
        
        for pattern in duration_patterns:
            match = re.search(pattern, text, re.IGNORECASE)
            if match:
                return match.group().strip()
        
        return ""
    
    def _extract_location(self, text: str) -> str:
        """Extract location from text."""
        # Common location patterns
        location_pattern = r'([A-Z][a-z]+,?\s*[A-Z][a-z]*,?\s*(India|IN)?)'
        match = re.search(location_pattern, text)
        return match.group().strip() if match else ""
    
    def _extract_institution(self, text: str) -> str:
        """Extract educational institution from text."""
        # Common institution indicators
        institution_keywords = ['University', 'College', 'Institute', 'School', 'IIT', 'IIM', 'NIT']
        
        for keyword in institution_keywords:
            if keyword.lower() in text.lower():
                # Extract surrounding text
                parts = text.split()
                for i, part in enumerate(parts):
                    if keyword.lower() in part.lower():
                        # Take surrounding words
                        start = max(0, i-2)
                        end = min(len(parts), i+3)
                        return ' '.join(parts[start:end])
        
        return ""
    
    def _extract_year(self, text: str) -> str:
        """Extract graduation year from text."""
        year_pattern = r'\b(19|20)\d{2}\b'
        match = re.search(year_pattern, text)
        return match.group() if match else ""
    
    def _extract_grade(self, text: str) -> str:
        """Extract grade/GPA from text."""
        grade_patterns = [
            r'(\d+\.?\d*)\s*(GPA|CGPA)',
            r'(\d+\.?\d*)%',
            r'(First Class|Second Class|Third Class|Distinction)',
        ]
        
        for pattern in grade_patterns:
            match = re.search(pattern, text, re.IGNORECASE)
            if match:
                return match.group().strip()
        
        return ""
    
    def _extract_certifications_enhanced(self, text: str) -> List[Dict[str, Any]]:
        """Enhanced certification extraction."""
        certifications = []
        
        # Look for certification section
        cert_keywords = ['certifications', 'certificates', 'credentials', 'licenses']
        sections = self._split_into_sections(text)
        
        for section_title, section_content in sections:
            if any(keyword in section_title.lower() for keyword in cert_keywords):
                # Parse individual certifications
                cert_entries = self._parse_certification_entries(section_content)
                certifications.extend(cert_entries)
        
        return certifications
    
    def _parse_certification_entries(self, content: str) -> List[Dict[str, Any]]:
        """Parse individual certification entries."""
        certifications = []
        
        lines = content.split('\n')
        for line in lines:
            line = line.strip()
            if not line or len(line) < 5:
                continue
            
            certification = {
                'name': line,
                'issuer': self._extract_issuer(line),
                'year': self._extract_year(line),
                'valid_until': self._extract_validity(line)
            }
            certifications.append(certification)
        
        return certifications
    
    def _extract_issuer(self, text: str) -> str:
        """Extract certification issuer."""
        # Common certification issuers
        issuers = ['Microsoft', 'Google', 'Amazon', 'Oracle', 'Cisco', 'IBM', 'Adobe', 'Salesforce']
        
        for issuer in issuers:
            if issuer.lower() in text.lower():
                return issuer
        
        return ""
    
    def _extract_validity(self, text: str) -> str:
        """Extract certification validity period."""
        validity_pattern = r'(Valid until|Expires?)\s*:?\s*(\w+ \d{4}|\d{4})'
        match = re.search(validity_pattern, text, re.IGNORECASE)
        return match.group(2) if match else ""
    
    def _extract_summary(self, text: str) -> str:
        """Extract professional summary/objective with better detection."""
        # Look for summary/objective sections
        summary_keywords = ['summary', 'objective', 'profile', 'about', 'overview', 'introduction']
        sections = self._split_into_sections(text)
        
        for section_title, section_content in sections:
            if any(keyword in section_title.lower() for keyword in summary_keywords):
                return section_content.strip()
        
        # Fallback: look for summary-like content in first few paragraphs
        lines = text.split('\n')
        for i, line in enumerate(lines[:10]):
            if any(keyword in line.lower() for keyword in summary_keywords):
                # Take next few lines as summary
                summary_lines = []
                for j in range(i+1, min(i+5, len(lines))):
                    if lines[j].strip():
                        summary_lines.append(lines[j].strip())
                    else:
                        break
                return ' '.join(summary_lines)
        
        return ""
    
    def _extract_additional_sections(self, text: str) -> Dict[str, Any]:
        """Extract additional sections like publications, volunteering, etc."""
        additional = {
            'publications': [],
            'volunteering': [],
            'hobbies': [],
            'references': [],
            'other': {}
        }
        
        sections = self._split_into_sections(text)
        
        for section_title, section_content in sections:
            title_lower = section_title.lower()
            
            if any(keyword in title_lower for keyword in ['publication', 'research', 'paper']):
                additional['publications'] = section_content.split('\n')
            elif any(keyword in title_lower for keyword in ['volunteer', 'community', 'social']):
                additional['volunteering'] = section_content.split('\n')
            elif any(keyword in title_lower for keyword in ['hobby', 'interest', 'personal']):
                additional['hobbies'] = section_content.split('\n')
            elif any(keyword in title_lower for keyword in ['reference', 'referee']):
                additional['references'] = section_content.split('\n')
            elif not any(keyword in title_lower for keyword in ['experience', 'education', 'skill', 'project', 'certification']):
                additional['other'][section_title] = section_content
        
        # Clean up empty lists
        return {k: v for k, v in additional.items() if v}
    
    def get_extraction_stats(self, file_path: str) -> Dict[str, Any]:
        """Get detailed statistics about the extraction process."""
        try:
            text = self.extract_text(file_path)
            structure = self.parse_resume_structure(text)
            
            stats = {
                'file_info': {
                    'path': file_path,
                    'size_bytes': os.path.getsize(file_path),
                    'detected_type': self.detect_file_type(file_path)
                },
                'extraction_stats': {
                    'total_characters': len(text),
                    'total_words': len(text.split()),
                    'total_lines': len(text.split('\n')),
                    'sections_found': len(structure.get('metadata', {}).get('sections_found', [])),
                    'contact_info_found': bool(structure.get('contact_info', {}).get('emails')),
                    'experience_entries': len(structure.get('experience', [])),
                    'education_entries': len(structure.get('education', [])),
                    'skills_found': sum(len(v) for v in structure.get('skills', {}).values() if isinstance(v, list)),
                    'certifications_found': len(structure.get('certifications', [])),
                    'projects_found': len(structure.get('projects', []))
                },
                'quality_indicators': {
                    'has_email': bool(structure.get('contact_info', {}).get('emails')),
                    'has_phone': bool(structure.get('contact_info', {}).get('phones')),
                    'has_experience': bool(structure.get('experience')),
                    'has_education': bool(structure.get('education')),
                    'has_skills': bool(structure.get('skills')),
                    'text_quality_score': self._calculate_text_quality_score(text)
                }
            }
            
            return stats
            
        except Exception as e:
            return {
                'error': str(e),
                'file_info': {
                    'path': file_path,
                    'size_bytes': os.path.getsize(file_path) if os.path.exists(file_path) else 0
                }
            }
    
    def _calculate_text_quality_score(self, text: str) -> float:
        """Calculate a quality score for extracted text (0-100)."""
        if not text:
            return 0.0
        
        score = 0.0
        
        # Length score (up to 20 points)
        if len(text) > 500:
            score += 20
        elif len(text) > 200:
            score += 15
        elif len(text) > 100:
            score += 10
        elif len(text) > 50:
            score += 5
        
        # Email presence (10 points)
        if re.search(r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b', text):
            score += 10
        
        # Phone presence (10 points)
        if re.search(r'(\+?\d{1,3}[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}', text):
            score += 10
        
        # Common resume sections (30 points total)
        section_keywords = ['experience', 'education', 'skills', 'work', 'employment', 'qualification']
        for keyword in section_keywords:
            if keyword in text.lower():
                score += 5
        
        # Professional terms (20 points)
        professional_terms = ['manager', 'developer', 'engineer', 'analyst', 'consultant', 'director', 'lead', 'senior']
        found_terms = sum(1 for term in professional_terms if term in text.lower())
        score += min(20, found_terms * 2)
        
        # Penalty for too many special characters (indicates extraction errors)
        special_char_ratio = len(re.findall(r'[^\w\s]', text)) / len(text)
        if special_char_ratio > 0.3:
            score -= 20
        elif special_char_ratio > 0.2:
            score -= 10
        
        return max(0.0, min(100.0, score))
    
    def validate_extraction(self, file_path: str) -> Dict[str, Any]:
        """Validate the quality of text extraction."""
        try:
            stats = self.get_extraction_stats(file_path)
            
            validation = {
                'is_valid': True,
                'confidence': 'high',
                'issues': [],
                'recommendations': []
            }
            
            quality_score = stats.get('quality_indicators', {}).get('text_quality_score', 0)
            
            if quality_score < 30:
                validation['is_valid'] = False
                validation['confidence'] = 'low'
                validation['issues'].append('Very low text quality score')
                validation['recommendations'].append('Consider re-scanning if this is an image-based document')
            elif quality_score < 60:
                validation['confidence'] = 'medium'
                validation['issues'].append('Moderate text quality')
                validation['recommendations'].append('Review extracted content for accuracy')
            
            # Check for essential elements
            if not stats.get('quality_indicators', {}).get('has_email'):
                validation['issues'].append('No email address found')
                validation['recommendations'].append('Verify contact information in original document')
            
            if not stats.get('quality_indicators', {}).get('has_experience'):
                validation['issues'].append('No work experience section found')
                validation['recommendations'].append('Check if experience is mentioned under different section names')
            
            if stats.get('extraction_stats', {}).get('total_words', 0) < 50:
                validation['is_valid'] = False
                validation['issues'].append('Very short extracted text')
                validation['recommendations'].append('Document may be corrupted or unsupported format')
            
            return validation
            
        except Exception as e:
            return {
                'is_valid': False,
                'confidence': 'none',
                'issues': [f'Validation failed: {str(e)}'],
                'recommendations': ['Check if file is accessible and in supported format']
            }
    
    def cleanup_file(self, file_path: str) -> bool:
        """
        Delete uploaded file.
        
        Args:
            file_path: Path to file to delete
            
        Returns:
            True if successful, False otherwise
        """
        try:
            if os.path.exists(file_path):
                os.remove(file_path)
                logger.info(f"File deleted: {file_path}")
                return True
            return False
        except Exception as e:
            logger.error(f"Failed to delete file {file_path}: {str(e)}")
            return False
